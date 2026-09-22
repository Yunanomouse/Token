#!/usr/bin/env python3
"""Structural and arithmetic checks on the 2026 dataset in data/json/.

Run with:  python3 -m unittest discover -s tests
"""
import csv
import importlib.util
import io
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
JSON_DIR = ROOT / "data" / "json"
CSV_DIR = ROOT / "data" / "csv"

JURISDICTIONS = ["AB", "BC", "SK", "MB", "ON", "QC", "NB", "NS", "PE", "NL", "YT", "NT", "NU"]


def load(name):
    with open(JSON_DIR / name, encoding="utf-8") as f:
        return json.load(f)


def cents(x):
    return round(x + 1e-9, 2)


class TestFileIntegrity(unittest.TestCase):
    def test_every_json_file_parses_and_is_for_2026(self):
        for path in sorted(JSON_DIR.glob("*.json")):
            with self.subTest(file=path.name):
                data = json.loads(path.read_text(encoding="utf-8"))
                year = data.get("tax_year") or data.get("as_of", "")[:4]
                self.assertIn("2026", str(year), f"{path.name} is not a 2026 dataset")


class TestPersonalIncomeTax(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data = load("personal_income_tax_2026.json")
        cls.jurs = {"CA": cls.data["federal"], **cls.data["provinces"]}

    def test_all_thirteen_provinces_and_territories_present(self):
        self.assertEqual(sorted(self.data["provinces"]), sorted(JURISDICTIONS))

    def test_brackets_are_ascending_and_open_ended(self):
        for code, jur in self.jurs.items():
            with self.subTest(jurisdiction=code):
                brackets = jur["brackets"]
                self.assertGreaterEqual(len(brackets), 2)
                self.assertIsNone(brackets[-1]["up_to"], "top bracket must be open-ended")
                uppers = [b["up_to"] for b in brackets[:-1]]
                self.assertNotIn(None, uppers, "only the top bracket may be open-ended")
                self.assertEqual(uppers, sorted(uppers), "thresholds must ascend")
                self.assertEqual(len(uppers), len(set(uppers)), "thresholds must be distinct")

    def test_marginal_rates_ascend_and_are_decimals(self):
        for code, jur in self.jurs.items():
            with self.subTest(jurisdiction=code):
                rates = [b["rate"] for b in jur["brackets"]]
                self.assertEqual(rates, sorted(rates), "marginal rates must ascend")
                for r in rates:
                    self.assertGreater(r, 0)
                    self.assertLess(r, 1, "rates are decimals, not percentages")

    def test_basic_personal_amounts_are_plausible(self):
        fed = self.data["federal"]["basic_personal_amount"]
        self.assertGreater(fed["max"], fed["min"])
        for code, jur in self.data["provinces"].items():
            bpa = jur["basic_personal_amount"]
            with self.subTest(jurisdiction=code):
                if bpa is None:
                    self.assertTrue(any("null" in n or "could not be verified" in n
                                        for n in jur.get("notes", [])),
                                    "a missing BPA must be explained in notes")
                elif isinstance(bpa, dict):
                    self.assertGreater(bpa["max"], bpa["min"])
                else:
                    self.assertGreater(bpa, 5000)
                    self.assertLess(bpa, 30000)

    def test_ontario_health_premium_bands_are_contiguous(self):
        bands = self.data["provinces"]["ON"]["health_premium"]["bands"]
        self.assertEqual(bands[0]["income_over"], 0)
        self.assertIsNone(bands[-1]["income_up_to"])
        for lower, upper in zip(bands, bands[1:]):
            self.assertEqual(lower["income_up_to"], upper["income_over"],
                             "bands must not leave a gap or overlap")
        self.assertEqual(bands[-1]["cap"], 900, "the premium is capped at $900")

    def test_ontario_surtax_thresholds_ascend(self):
        thresholds = self.data["provinces"]["ON"]["surtax"]["thresholds"]
        self.assertEqual([t["tax_over"] for t in thresholds],
                         sorted(t["tax_over"] for t in thresholds))


class TestPayrollContributions(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.pay = load("payroll_contributions_2026.json")

    def test_base_plan_maximums_match_rate_times_pensionable_earnings(self):
        for plan in ("cpp", "qpp"):
            p = self.pay[plan]
            pensionable = p["ympe"] - p["basic_exemption"]
            with self.subTest(plan=plan):
                self.assertEqual(cents(pensionable * p["employee_rate"]),
                                 p["max_employee_contribution"])
                self.assertEqual(cents(pensionable * p["self_employed_rate"]),
                                 p["max_self_employed_contribution"])
                self.assertEqual(cents(p["employee_rate"] * 2), cents(p["self_employed_rate"]))

    def test_second_tier_maximums_match_the_ympe_to_yampe_band(self):
        for plan in ("cpp2", "qpp2"):
            p = self.pay[plan]
            band = p["earnings_band"]
            with self.subTest(plan=plan):
                self.assertEqual(band["to"], p["yampe"])
                width = band["to"] - band["from"]
                self.assertEqual(cents(width * p["employee_rate"]), p["max_employee_contribution"])
                self.assertEqual(cents(width * p["self_employed_rate"]),
                                 p["max_self_employed_contribution"])

    def test_second_tier_starts_where_the_base_plan_ends(self):
        self.assertEqual(self.pay["cpp2"]["earnings_band"]["from"], self.pay["cpp"]["ympe"])
        self.assertEqual(self.pay["qpp2"]["earnings_band"]["from"], self.pay["qpp"]["ympe"])

    def test_ei_and_qpip_maximums_match_rate_times_insurable_earnings(self):
        ei = self.pay["ei"]
        mie = ei["maximum_insurable_earnings"]
        self.assertEqual(cents(mie * ei["employee_rate"]), ei["max_employee_premium"])
        self.assertEqual(cents(ei["employee_rate"] * ei["employer_rate_multiplier"]),
                         cents(ei["employer_rate"]))
        self.assertEqual(cents(mie * ei["quebec"]["employee_rate"]),
                         ei["quebec"]["max_employee_premium"])
        self.assertLess(ei["quebec"]["employee_rate"], ei["employee_rate"],
                        "Quebec's EI rate is reduced because QPIP covers parental benefits")

        qpip = self.pay["qpip"]
        qmie = qpip["maximum_insurable_earnings"]
        self.assertEqual(cents(qmie * qpip["employee_rate"]), qpip["max_employee_premium"])
        self.assertEqual(cents(qmie * qpip["employer_rate"]),
                         qpip["max_employer_premium_per_employee"])


class TestSalesTax(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data = load("sales_tax_2026.json")

    def test_all_thirteen_jurisdictions_present(self):
        self.assertEqual(sorted(self.data["jurisdictions"]), sorted(JURISDICTIONS))

    def test_total_rate_is_the_sum_of_its_components(self):
        for code, j in self.data["jurisdictions"].items():
            with self.subTest(jurisdiction=code):
                parts = (j["hst"],) if j["hst"] is not None else (j["gst"], j["pst"])
                self.assertEqual(cents(sum(p for p in parts if p)), cents(j["total_rate"]))

    def test_gst_rate_is_consistent_in_non_harmonized_provinces(self):
        for code, j in self.data["jurisdictions"].items():
            if j["hst"] is None:
                with self.subTest(jurisdiction=code):
                    self.assertEqual(j["gst"], self.data["federal_gst_rate"])


class TestCorporateIncomeTax(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data = load("corporate_income_tax_2026.json")

    def test_all_thirteen_provinces_present(self):
        self.assertEqual(sorted(self.data["provincial"]), sorted(JURISDICTIONS))

    def test_small_business_rate_never_exceeds_the_general_rate(self):
        for code, j in {"CA": self.data["federal"], **self.data["provincial"]}.items():
            with self.subTest(jurisdiction=code):
                self.assertLessEqual(j["small_business_rate"], j["general_rate"])
                self.assertGreaterEqual(j["small_business_limit"], 500000)


class TestCreditsAndLimits(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data = load("credits_and_limits_2026.json")

    def test_tfsa_cumulative_room_matches_the_annual_limits_since_2009(self):
        annual = ([5000] * 4 + [5500] * 2 + [10000] + [5500] * 3
                  + [6000] * 4 + [6500] + [7000] * 3)  # 2009-2026
        tfsa = self.data["registered_accounts"]["tfsa"]
        self.assertEqual(len(annual), 2026 - 2009 + 1)
        self.assertEqual(annual[-1], tfsa["annual_limit"])
        self.assertEqual(sum(annual), tfsa["cumulative_room_since_2009"])

    def test_fhsa_single_year_maximum_is_the_annual_limit_plus_carry_forward(self):
        fhsa = self.data["registered_accounts"]["fhsa"]
        self.assertIn(f"${fhsa['annual_limit'] + fhsa['max_carry_forward']:,}", fhsa["note"])
        self.assertLessEqual(fhsa["annual_limit"], fhsa["lifetime_limit"])

    def test_federal_credit_rate_matches_the_lowest_federal_bracket(self):
        fed = load("personal_income_tax_2026.json")["federal"]
        self.assertEqual(self.data["credits"]["credit_rate_federal"], fed["brackets"][0]["rate"])

    def test_federal_bpa_matches_the_personal_income_tax_file(self):
        fed = load("personal_income_tax_2026.json")["federal"]["basic_personal_amount"]
        credits = self.data["credits"]["federal_basic_personal_amount"]
        self.assertEqual(credits["max"], fed["max"])
        self.assertEqual(credits["min"], fed["min"])

    def test_capital_gains_inclusion_rate_is_one_half(self):
        self.assertEqual(self.data["capital_gains"]["inclusion_rate"], 0.5)


class TestCsvExports(unittest.TestCase):
    """The committed CSVs must match what generate_csv.py produces today."""

    @classmethod
    def setUpClass(cls):
        spec = importlib.util.spec_from_file_location(
            "generate_csv", ROOT / "scripts" / "generate_csv.py")
        cls.gen = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(cls.gen)

    def test_committed_csvs_are_up_to_date(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.gen.CSV_DIR = Path(tmp)
            self.gen.ROOT = Path(tmp)  # keeps the script's progress printing relative
            stdout, sys.stdout = sys.stdout, io.StringIO()
            try:
                self.gen.brackets_csv()
                self.gen.bpa_csv()
                self.gen.sales_tax_csv()
                self.gen.corporate_csv()
                self.gen.payroll_csv()
            finally:
                sys.stdout = stdout

            regenerated = sorted(Path(tmp).glob("*.csv"))
            self.assertEqual([p.name for p in regenerated],
                             sorted(p.name for p in CSV_DIR.glob("*.csv")),
                             "the set of CSV exports changed")
            for path in regenerated:
                with self.subTest(file=path.name):
                    self.assertEqual(
                        (CSV_DIR / path.name).read_text(encoding="utf-8"),
                        path.read_text(encoding="utf-8"),
                        f"data/csv/{path.name} is stale — run python3 scripts/generate_csv.py")

    def test_bracket_csv_rows_cover_every_jurisdiction(self):
        with open(CSV_DIR / "income_tax_brackets_2026.csv", encoding="utf-8") as f:
            codes = {row["jurisdiction_code"] for row in csv.DictReader(f)}
        self.assertEqual(codes, set(JURISDICTIONS) | {"CA"})


class TestWorkbookBuilder(unittest.TestCase):
    """scripts/build_tax_workbook.py must still understand the dataset's shapes."""

    def test_workbook_builds_from_the_current_dataset(self):
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            self.skipTest("openpyxl is not installed (pip install openpyxl)")

        import os
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "workbook.xlsx"
            result = subprocess.run(
                [sys.executable, str(ROOT / "scripts" / "build_tax_workbook.py")],
                capture_output=True, text=True,
                env={**os.environ, "TAX_WORKBOOK_OUT": str(out)})
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertTrue(out.is_file(), "no workbook was written")

            from openpyxl import load_workbook
            sheets = load_workbook(out).sheetnames
            for expected in ("Inputs", "Tax Data 2026", "Payroll CPP EI"):
                self.assertIn(expected, sheets)


if __name__ == "__main__":
    unittest.main()
